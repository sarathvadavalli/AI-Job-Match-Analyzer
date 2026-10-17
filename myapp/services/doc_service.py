from io import BytesIO
import os
import pymupdf
import pytesseract
from PIL import Image
from dotenv import load_dotenv
from docx import Document
from zipfile import ZipFile, BadZipFile

from myapp.schemas.resume_extraction import ResumeExtractionInput
import logging

logger = logging.getLogger()

load_dotenv()
pytesseract.pytesseract.tesseract_cmd = os.getenv("TESSERACT_CMD") or "tesseract"

MAX_UPLOAD_BYTES = 4 * 1024 * 1024
MAX_DOCUMENT_CHARS = 100_000
MAX_PDF_PAGES = 20
MAX_OCR_PAGES = 5

class DocumentService:

    def __init__(self, file_type):
        self.file_type = file_type
 
    def extract_text(self, file_name: str, content: bytes) -> ResumeExtractionInput:
        if not content:
            raise ValueError("The uploaded file is empty.")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("Upload a document smaller than 4 MiB.")
        extension = file_name.lower().split(".")[-1]

        if extension == "pdf":
            result = self._extract_pdf(content)

        elif extension == "docx":
            result = self._extract_docx(content)

        elif extension == "txt":
            result = self._extract_txt(content)
        else:
            raise ValueError("Unsupported file type")

        if not result.text.strip():
            raise ValueError("No readable text was found in the document.")
        if len(result.text) > MAX_DOCUMENT_CHARS:
            raise ValueError("Document is too long. Upload at most 100,000 characters.")
        return result

    def _extract_pdf(self, content: bytes) -> ResumeExtractionInput:
        text = []
        links = []

        try:
            with pymupdf.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise ValueError("Upload an unencrypted PDF.")
                if len(document) > MAX_PDF_PAGES:
                    raise ValueError("Upload a PDF with at most 20 pages.")
                scanned_pages = [page.number for page in document if not page.get_text().strip()]
                if len(scanned_pages) > MAX_OCR_PAGES:
                    raise ValueError("Upload a PDF with at most 5 scanned pages.")

                for page in document:
                    page_text = page.get_text().strip()
                    if page.number in scanned_pages:
                        page_text = self._extract_page_ocr(page)
                    page_links = [
                        link["uri"]
                        for link in page.get_links()
                        if "uri" in link
                    ]
                    links.extend(page_links)
                    text.append("\n".join([page_text, *page_links]))
                    if sum(map(len, text)) > MAX_DOCUMENT_CHARS:
                        raise ValueError("Document is too long. Upload at most 100,000 characters.")

            return ResumeExtractionInput(text="\n".join(text).strip(), links=list(dict.fromkeys(links)))
        except (pymupdf.FileDataError, pymupdf.EmptyFileError, pymupdf.mupdf.FzErrorBase) as exc:
            raise ValueError("Unable to read this PDF. Upload a valid PDF file.") from exc

    def _extract_page_ocr(self, page) -> str:
        """Recognize one bounded grayscale page locally with Tesseract."""
        longest_side = max(page.rect.width, page.rect.height)
        if longest_side <= 0:
            raise ValueError("The PDF contains a page with invalid dimensions.")
        scale = min(150 / 72, 1600 / longest_side)
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csGRAY, alpha=False)
        try:
            with Image.frombytes("L", (pixmap.width, pixmap.height), pixmap.samples) as image:
                del pixmap
                return pytesseract.image_to_string(image, lang="eng", timeout=30).strip()
        except pytesseract.TesseractNotFoundError as e:
            raise RuntimeError("Tesseract OCR is not installed or TESSERACT_CMD is invalid.") from e
        except Exception as e:
            logger.exception("Tesseract OCR failed")
            raise RuntimeError("Scanned PDF recognition failed. Try again or upload a text-based PDF.") from e


    def _extract_docx(self, content: bytes) -> ResumeExtractionInput:
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("DOCX content is too large after decompression.")
        except BadZipFile as exc:
            raise ValueError("Upload a valid DOCX file.") from exc
        document = Document(BytesIO(content))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        links = []
        for paragraph in document.paragraphs:
            if self.file_type == "resume":
                for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
                    rel_id = hyperlink.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )

                    if rel_id:
                        relationship = paragraph.part.rels.get(rel_id)

                        if relationship and relationship.target_ref:
                            links.append(relationship.target_ref)

        text = "\n".join(paragraphs).strip()
        links = list(dict.fromkeys(links))
        return ResumeExtractionInput(text=text, links=links)

    def _extract_txt(self, content: bytes) -> ResumeExtractionInput:
        try:
            return ResumeExtractionInput(text=content.decode("utf-8-sig").strip(), links=[])
        except UnicodeDecodeError as exc:
            raise ValueError("TXT documents must use UTF-8 encoding.") from exc


# try:
#     file_path = Path("D:/AI Projects/JD Extractor/Sample docs/LL.pptx")
#     if not file_path.exists():
#         raise FileNotFoundError(f"File not found: {file_path}")

#     content = file_path.read_bytes()

#     text = DocumentService.extract_text(
#         file_name=file_path.name,
#         content=content
#     )
#     return text
# except Exception as e:
#     return f"Error: {str(e)}"
