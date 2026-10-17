"""Offline regression checks for serverless document processing.

Run with: python -m unittest tests.test_vercel
"""
from io import BytesIO
import os
import unittest
from unittest.mock import patch

import pymupdf
from docx import Document

from myapp.services.doc_service import DocumentService, MAX_UPLOAD_BYTES


def make_pdf(texts):
    with pymupdf.open() as document:
        for text in texts:
            page = document.new_page()
            if text:
                page.insert_text((72, 72), text)
        return document.tobytes()


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.service = DocumentService("resume")

    def test_text_pdf_does_not_use_ocr_and_keeps_links(self):
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), "Candidate resume")
            page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 72, 180, 90), "uri": "https://example.com"})
            content = document.tobytes()
        with patch("myapp.services.doc_service.pytesseract.image_to_string") as ocr:
            result = self.service.extract_text("resume.pdf", content)
        ocr.assert_not_called()
        self.assertIn("Candidate resume", result.text)
        self.assertEqual(result.links, ["https://example.com"])

    def test_mixed_pdf_only_sends_scanned_page_to_tesseract(self):
        with patch("myapp.services.doc_service.pytesseract.image_to_string", return_value="Scanned experience") as ocr:
            result = self.service.extract_text("resume.pdf", make_pdf(["Digital text", ""]))
            image = ocr.call_args.args[0]
            self.assertEqual(image.mode, "L")
            self.assertLessEqual(max(image.size), 1601)
            ocr.assert_called_once_with(image, lang="eng", timeout=30)
        self.assertIn("Digital text", result.text)
        self.assertIn("Scanned experience", result.text)

    def test_page_limits_checked_before_ocr(self):
        with patch("myapp.services.doc_service.pytesseract.image_to_string") as ocr:
            for pages, message in [([""] * 6, "5 scanned"), (["text"] * 21, "20 pages")]:
                with self.assertRaisesRegex(ValueError, message):
                    self.service.extract_text("resume.pdf", make_pdf(pages))
        ocr.assert_not_called()

    def test_invalid_and_oversized_documents(self):
        for name, data in [("x.pdf", b"invalid"), ("x.txt", b""), ("x.txt", b"x" * (MAX_UPLOAD_BYTES + 1)), ("x.docx", b"invalid")]:
            with self.subTest(name=name, size=len(data)):
                with self.assertRaises(ValueError):
                    self.service.extract_text(name, data)
                    
    def test_docx_paragraphs_are_not_duplicated(self):
        document = Document()
        document.add_paragraph("Candidate experience")
        content = BytesIO()
        document.save(content)
        self.assertEqual(self.service.extract_text("x.docx", content.getvalue()).text, "Candidate experience")

    def test_missing_tesseract_has_actionable_error(self):
        import pytesseract
        with patch("myapp.services.doc_service.pytesseract.image_to_string", side_effect=pytesseract.TesseractNotFoundError()):
            with self.assertRaisesRegex(RuntimeError, "Tesseract OCR is not installed"):
                self.service.extract_text("scan.pdf", make_pdf([""]))


class RouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Use explicit placeholders so the offline test never connects to real services.
        cls.environment = patch.dict(os.environ, {
            "MONGODB_URI": "mongodb://localhost:27017",
            "JWT_SECRET_KEY": "offline-test-secret",
            "OPENROUTER_API_KEY": "offline-test-key",
        })
        cls.environment.start()
        from app import app
        from fastapi.testclient import TestClient
        cls.app = app
        cls.test_client = TestClient

    @classmethod
    def tearDownClass(cls):
        cls.environment.stop()

    def test_pages_upload_errors_and_resume_response(self):
        from myapp.schemas.resume_extraction import ResumeProfile
        with patch("myapp.main.init_db"), patch("myapp.main.close_db"), self.test_client(self.app) as client:
            self.assertEqual(client.get("/login").status_code, 200)
            self.assertEqual(client.get("/dashboard", headers={"Accept": "text/html"}, follow_redirects=False).status_code, 303)
            self.assertEqual(client.post("/resume/extract", files={"file": ("x.pdf", b"bad", "application/pdf")}).status_code, 400)
            self.assertEqual(client.post("/resume/extract", files={"file": ("x.txt", b"x" * (MAX_UPLOAD_BYTES + 1), "text/plain")}).status_code, 413)
            with patch("myapp.routes.document_routes.ResumeService") as service:
                service.return_value.extract_resume_info.return_value = ResumeProfile(name="Candidate")
                response = client.post("/resume/extract", files={"file": ("x.txt", b"Candidate", "text/plain")})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["name"], "Candidate")


if __name__ == "__main__":
    unittest.main()
